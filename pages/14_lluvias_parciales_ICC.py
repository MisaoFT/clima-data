import io
import pandas as pd
import numpy as np
import streamlit as st
from openpyxl import Workbook

# Títulos de la interfaz web para esta herramienta específica
st.title("📊 Procesador de lluvias parciales obtenidas de ICC")
st.write("Formatos: Inyección matricial directa con fila de códigos e identificación adaptativa de columnas.")
st.markdown("---")

# 1. COMPONENTE WEB PARA SUBIR EL ARCHIVO DEL ICC (Acepta CSV y Excel)
st.subheader("1. Carga de archivo obtenido de ICC")
subido_icc = st.file_uploader(
    "Por favor, selecciona tu archivo de datos extraído del ICC (CSV o Excel):", 
    type=["csv", "xlsx", "xls"]
)

if subido_icc is not None:
    nombre_archivo = subido_icc.name
    st.info(f"Archivo seleccionado: **{nombre_archivo}**")
    
    # 2. LECTURA ADAPTATIVA DEL ARCHIVO
    try:
        if nombre_archivo.endswith('.csv'):
            df_raw = pd.read_csv(subido_icc, sep=None, engine='python')
        else:
            df_raw = pd.read_excel(subido_icc)
        st.success(f"¡Archivo del ICC leído correctamente! ({len(df_raw)} registros detectados).")
        
        # --- PROCESAMIENTO INTERNO EXACTAMENTE IGUAL ---
        # Limpieza y estándar de columnas
        df_raw.columns = [col.lower().strip() for col in df_raw.columns]

        # Identificación automática de columnas críticas
        col_estacion = next((c for c in df_raw.columns if 'estacion' in c or 'nombre' in c or 'station' in c), None)
        col_fecha = next((c for c in df_raw.columns if 'fecha' in c or 'date' in c), None)
        col_lluvia = next((c for c in df_raw.columns if 'lluvia' in c or 'precip' in c or 'pp' in c or 'val' in c), None)

        if not col_estacion or not col_fecha or not col_lluvia:
            # Forzado por posición si los nombres varían mucho
            df_raw['Estacion_Limpia'] = df_raw.iloc[:, 0].astype(str).str.replace('_', ' ').str.upper().str.strip()
            df_raw['Fecha_Limpia'] = pd.to_datetime(df_raw.iloc[:, 1], errors='coerce').dt.strftime('%Y-%m-%d')
            df_raw['Valor_Lluvia'] = pd.to_numeric(df_raw.iloc[:, 2], errors='coerce')
        else:
            df_raw['Estacion_Limpia'] = df_raw[col_estacion].astype(str).str.replace('_', ' ').str.upper().str.strip()
            df_raw['Fecha_Limpia'] = pd.to_datetime(df_raw[col_fecha], errors='coerce').dt.strftime('%Y-%m-%d')
            df_raw['Valor_Lluvia'] = pd.to_numeric(df_raw[col_lluvia], errors='coerce')

        # Consolidar lluvia diaria total por estación
        df_resumen = df_raw.groupby(['Fecha_Limpia', 'Estacion_Limpia'], as_index=False)['Valor_Lluvia'].sum()

        # 3. DICCIONARIO ESTRUCTURAL (Las 29 estaciones del ICC)
        estaciones_icc = {
            "ALAMO": "ICC121701AT", "AMAZONAS": "ICC050501AT", "BONANZA": "ICC050701AT",
            "BOUGANVILIA": "ICC050301AT", "CENGICAÑA": "ICC050202AT", "CHIQUIRINES": "ICC110101AT",
            "COCALES": "ICC101401AT", "CONCEPCIÓN": "ICC050101AT", "COSTA BRAVA": "ICC050302AT",
            "EL BÁLSAMO": "ICC050203AT", "EL PLATANAR": "ICC041101AT", "IRLANDA": "ICC050601AT",
            "LA CANDELARIA": "ICC060901AT", "LA GIRALDA": "ICC050901AT", "LA MAQUINA": "ICC060801AT",
            "LORENA": "ICC101001AT", "NARANJALES": "ICC100601AT", "PETÉN OFICINA": "ICC050602AT",
            "PROVIDENCIA": "ICC110701AT", "PUYUMATE": "ICC051301AT", "SAN ANTONIO EV": "ICC051401AT",
            "SAN NICOLÁS": "ICC100701AT", "SAN RAFAEL": "ICC060902AT", "TEHUANTEPEQ": "ICC050201AT",
            "TRINIDAD": "ICC050502AT", "TRINIDAD MAGDALENA": "ICC221501AT", "TULULÁ": "ICC110601AT",
            "XOLUTA": "ICC110102AT", "YEPOCAPA (FCA-CATIE)": "ICC041201AT"
        }

        # 4. CREACIÓN DEL EXCEL DESDE CERO
        wb = Workbook()
        ws = wb.active
        ws.title = "Hoja1"

        ws.cell(row=1, column=1, value="FECHA")
        ws.cell(row=1, column=2, value="VARIABLE")

        mapeo_columnas = {}
        col_idx = 3 
        for nombre, codigo in estaciones_icc.items():
            ws.cell(row=1, column=col_idx, value=codigo) 
            mapeo_columnas[nombre] = col_idx
            col_idx += 1

        # 5. ORDENAMIENTO CRONOLÓGICO
        fechas_unicas = sorted(df_resumen['Fecha_Limpia'].dropna().unique())

        # 6. INYECCIÓN MATRICIAL DE DATOS
        fila_actual = 2
        for fecha_v in fechas_unicas:
            ws.cell(row=fila_actual, column=1, value=fecha_v)
            ws.cell(row=fila_actual, column=2, value="PREP")

            # Inicializar con "NA"
            for c_idx in mapeo_columnas.values():
                ws.cell(row=fila_actual, column=c_idx, value="NA")

            datos_dia = df_resumen[df_resumen['Fecha_Limpia'] == fecha_v]

            for _, r_dato in datos_dia.iterrows():
                estacion_raw = r_dato['Estacion_Limpia']
                valor_lluvia = r_dato['Valor_Lluvia']

                estacion_match = None
                for nombre_estandar in mapeo_columnas.keys():
                    if nombre_estandar in estacion_raw or estacion_raw in nombre_estandar:
                        estacion_match = nombre_estandar
                        break

                if estacion_match:
                    col_destino = mapeo_columnas[estacion_match]
                    if pd.notna(valor_lluvia):
                        ws.cell(row=fila_actual, column=col_destino, value=valor_lluvia)

            fila_actual += 1

        # Muestra una vista previa de los datos consolidados en pantalla
        st.markdown("---")
        st.subheader("2. Vista previa de la data")
        st.dataframe(df_resumen.head(10), use_container_width=True)

        # 7. GUARDAR EN MEMORIA PARA DESCARGA
        buffer_excel = io.BytesIO()
        wb.save(buffer_excel)
        buffer_excel.seek(0)

        # BOTÓN WEB PARA DESCARGAR EL RESULTADO
        st.markdown("---")
        st.subheader("3. Descargar archivo generado")
        
        st.download_button(
            label="💾 Descargar reporte de lluvias parciales de estaciones ICC (.xlsx)",
            data=buffer_excel,
            file_name="datos_lluvias_parciales_ICC_fecha_hora.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
    except Exception as e:
        st.error(f"Ocurrió un error al procesar el archivo del ICC: {e}")

else:
    st.info("A la espera de que subas un archivo CSV o Excel para iniciar el análisis.")