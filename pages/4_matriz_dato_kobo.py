import io
import pandas as pd
import numpy as np
import streamlit as st
from openpyxl import Workbook
from datetime import datetime

# Títulos y estilos visuales en la interfaz
st.title("🌧️ Procesador de lluvias parciales obtenidas de Kobo")
st.write("Formatos: Estructura fija cronológica con relleno 'NA' y mapeo directo de estaciones.")
st.markdown("---")

# 1. COMPONENTE WEB PARA SUBIR EL ARCHIVO DE KOBO
st.subheader("1. Carga de archivo de KoboCollect")
subido_kobo = st.file_uploader(
    "Por favor, selecciona tu archivo CSV descargado de Kobo:", 
    type=["csv"]
)

if subido_kobo is not None:
    nombre_csv = subido_kobo.name
    st.info(f"Archivo seleccionado: **{nombre_csv}**")
            
    # 2. PROCESAMIENTO INTERNO DEL CSV DE KOBO
    try:
        df_kobo = pd.read_csv(subido_kobo, sep=';', engine='python')
        st.success(f"¡Archivo de Kobo leído correctamente! ({len(df_kobo)} registros detectados).")
                        
        # --- PROCESAMIENTO DE DATOS ---
        df_kobo['Estacion_Limpia'] = None
        for col in ['estacion_auto', 'estacion_conven', 'estacion_sinop']:
            if col in df_kobo.columns:
                df_kobo['Estacion_Limpia'] = df_kobo['Estacion_Limpia'].fillna(df_kobo[col])
                
        # Limpieza estándar básica (Cambiamos guiones bajos por espacios para que se parezcan a los nombres reales)
        df_kobo['Estacion_Limpia'] = df_kobo['Estacion_Limpia'].fillna("DESCONOCIDA").astype(str).str.replace('_', ' ').str.upper().str.strip()
                
        # Normalización de fechas e horarios
        df_kobo['Fecha_Limpia'] = pd.to_datetime(df_kobo['fecha'], errors='coerce').dt.strftime('%Y-%m-%d')
        df_kobo['Hora_Limpia'] = df_kobo['Hora'].str.replace('_', ':') 
                
        # Extracción numérica de lluvias
        df_kobo['lluvia_13'] = pd.to_numeric(df_kobo['lluvia_13'], errors='coerce')
        df_kobo['lluvia_18'] = pd.to_numeric(df_kobo['lluvia_18'], errors='coerce')
                
        df_kobo['Valor_Lluvia'] = np.where(df_kobo['Hora_Limpia'] == '13:00', df_kobo['lluvia_13'], df_kobo['lluvia_18'])
                
        # --- CONTROL DE DUPLICADOS EN LA MATRIZ (EVITA SUMATORIAS) ---
        # Si un observador envió dos reportes para el mismo día y hora, 
        # conservamos únicamente el último registro disponible en la base de datos de Kobo.
        df_kobo = df_kobo.sort_index(ascending=True) # Mantiene el orden de inserción de filas
        df_resumen = df_kobo.drop_duplicates(
            subset=['Fecha_Limpia', 'Hora_Limpia', 'Estacion_Limpia'], 
            keep='last'
        )[['Fecha_Limpia', 'Hora_Limpia', 'Estacion_Limpia', 'Valor_Lluvia']]
                
        # 3. DICCIONARIO ESTRUCTURAL INTEGRADO (Las 64 estaciones oficiales)
        estaciones_estructura = {
            "CAMOTAN": "INS200501CV", "CATARINA": "INS121601CV", "CHAMPERICO FEGUA": "INS110701CV",
            "COBAN": "INS160101CV", "EL CAPITAN": "INS071301CV", "ESQUIPULAS": "INS200701CV",
            "FLORES AEROPUERTO": "INS170101CV", "LA AURORA": "INS010102CV", "LA CEIBITA": "INS210601CV",
            "LA FRAGUA": "INS190201CV", "LABOR OVALLE": "INS090301CV", "LOS ALTOS": "INS090101CV",
            "MAZATENANGO": "INS100101CV", "MONTUFAR": "INS221401CV", "PANZÓS PHC ALTAVERAPAZ": "INS160701CV",
            "PUERTO BARRIOS PHC": "INS180101CV", "RETALHULEU AEROPUERTO": "INS110101CV", "SACAPULAS": "INS141601CV",
            "SAN MARCOS PHC": "INS120101CV", "SANTA CRUZ BALANYÁ": "INS041001CV", "TODOS SANTOS": "INS131501CV",
            "INSIVUMEH": "INS010101CV", "ALAMEDA ICTA": "INS040101CV", "ASUNCION MITA": "INS220501CV",
            "CHINIQUE": "INS140301CV", "EL TABLON": "INS070101CV", "HUEHUETENANGO": "INS130101CV",
            "LA UNION": "INS190901CV", "LAS VEGAS PHC": "INS180201CV", "NEBAJ": "INS141301CV",
            "POPTUN": "INS171201CV", "QUEZADA": "INS221701CV", "SABANA GRANDE": "INS050101CV",
            "SAN AGUSTIN ACASAGUASTLAN": "INS020302CV", "SAN JERONIMO": "INS150701CV", "SAN JOSE AEROPUERTO": "INS050901CV",
            "SAN PEDRO NECTA": "INS130601CV", "SANTA MARIA CAHABON": "INS161201CV", "SANTIAGO ATITLAN": "INS071901CV",
            "SUIZA CONTENTA": "INS030801CV", "TECUN UMAN": "INS121701CV", "CHIXOY PCH": "INS141901CV",
            "CUBULCO": "INS150301CV", "LOS ALBORES": "INS150101CV", "LOS ESCLAVOS": "INS060201CV",
            "PASABIEN": "INS190301CV", "POTRERO CARRILLO": "INS020501CV", "SAN MARTIN JILOTEPEQUE": "INS040501CV",
            "AMATITLAN": "INS011401AT", "ANTIGUA GUATEMALA": "INS030101CV", "CONCEPCION": "INS120701CV",
            "IXCHIGUAN": "INS121001CV", "JALAPA": "INS210101CV", "LA REFORMA": "INS121301CV",
            "LO DE COY": "INS010701CV", "MARISCOS": "INS180501AT", "MORALES MET": "INS180401AT",
            "PACHUTÉ": "INS040701CV", "PLAYA GRANDE IXCAN": "INS142101CV", "SAN JOSE PINULA": "INS011001CV",
            "SAN PEDRO AYAMPUC": "INS010401CV", "SANTA CRUZ DEL QUICHE": "INS140101CV", "SANTA MARGARITA": "INS190501CV",
            "TOTONICAPAN": "INS080101CV"
        }
                
        # 4. TRADUCTOR ADAPTATIVO PARA ERRORES CRÍTICOS DE KOBO
        traductor_kobo = {
            "PANZ S PHC ALTAVERAPAZ": "PANZÓS PHC ALTAVERAPAZ",
            "ASUNCI N MITA": "ASUNCION MITA",
            "SANTA CRUZ BALANY": "SANTA CRUZ BALANYÁ",
            "PACHUTE": "PACHUTÉ",
            "CONCEPCI N": "CONCEPCION"
        }
                
        # 5. CREACIÓN DEL EXCEL EN MEMORIA
        wb = Workbook()
        ws = wb.active
        ws.title = "Hoja1"
        ws.cell(row=2, column=1, value="FECHA")
        ws.cell(row=2, column=2, value="HORA")
        ws.cell(row=2, column=3, value="VARIABLE")
                
        mapeo_columnas = {}
        col_idx = 4
        for nombre, codigo in estaciones_estructura.items():
            ws.cell(row=1, column=col_idx, value=nombre)
            ws.cell(row=2, column=col_idx, value=codigo)
            mapeo_columnas[nombre] = col_idx
            col_idx += 1
                    
        # 6. ORDENAMIENTO CRONOLÓGICO
        fechas_horas_unicas = df_resumen[['Fecha_Limpia', 'Hora_Limpia']].drop_duplicates()
        fechas_horas_unicas = fechas_horas_unicas.sort_values(by=['Fecha_Limpia', 'Hora_Limpia'], ascending=[True, True])
                
        # 7. INYECCIÓN MATRICIAL DIRECTA Y SEGURA
        fila_actual = 3
        for _, r_tiempo in fechas_horas_unicas.iterrows():
            fecha_v = r_tiempo['Fecha_Limpia']
            hora_v = r_tiempo['Hora_Limpia']
            ws.cell(row=fila_actual, column=1, value=fecha_v)
            ws.cell(row=fila_actual, column=2, value=f"{hora_v} HORAS")
            ws.cell(row=fila_actual, column=3, value="PREP")
                        
            # Inicializar con "NA" todas las columnas de la fila
            for c_idx in mapeo_columnas.values():
                ws.cell(row=fila_actual, column=c_idx, value="NA")
                            
            datos_bloque = df_resumen[(df_resumen['Fecha_Limpia'] == fecha_v) & (df_resumen['Hora_Limpia'] == hora_v)]
                        
            for _, r_dato in datos_bloque.iterrows():
                estacion_kobo = r_dato['Estacion_Limpia']  
                valor_lluvia = r_dato['Valor_Lluvia']
                                
                if estacion_kobo in traductor_kobo:
                    estacion_match = traductor_kobo[estacion_kobo]
                elif estacion_kobo in mapeo_columnas:
                    estacion_match = estacion_kobo
                else:
                    estacion_match = None
                                        
                if estacion_match and estacion_match in mapeo_columnas:
                    col_destino = mapeo_columnas[estacion_match]
                    if pd.notna(valor_lluvia):
                        ws.cell(row=fila_actual, column=col_destino, value=valor_lluvia)
            
            fila_actual += 1
            
        # Muestra una vista previa en la web
        st.markdown("---")
        st.subheader("2. Vista previa de la data procesada")
        st.dataframe(df_resumen.head(10), use_container_width=True)
                
        # 8. GUARDAR EN MEMORIA PARA DESCARGA WEB
        buffer_excel = io.BytesIO()
        wb.save(buffer_excel)
        buffer_excel.seek(0)
                
        ahora_ejecucion = datetime.now()
        nombre_salida_dinamico = f"datos-diarios_prelim_prep-milim_parcial_{ahora_ejecucion.strftime('%Y%m%d_%H%M')}.xlsx"
                
        # BOTÓN WEB PARA DESCARGAR
        st.markdown("---")
        st.subheader("3. Descargar archivo generado")
                        
        st.download_button(
            label="💾 Descargar documento de lluvias parciales procesado (.xlsx)",
            data=buffer_excel,
            file_name=nombre_salida_dinamico,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
                
        # --- LIMPIEZA DE MEMORIA POST-PROCESAMIENTO (KOBO) ---
        del df_kobo
        del df_resumen
        del datos_bloque
                
        import gc
        gc.collect()
                
    except Exception as e:
        st.error(f"Ocurrió un error al procesar el archivo CSV: {e}")
else:
    st.info("A la espera de que subas un archivo CSV para iniciar el análisis.")